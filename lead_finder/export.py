"""CSV / Excel exports with friendly column names."""
from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .contacts import readiness_score

EXPORT_COLUMNS = {
    "business_name": "Business", "category": "Category", "business_phone": "Phone",
    "decision_maker": "Decision-maker", "decision_maker_title": "Title", "decision_maker_confidence": "Confidence",
    "email": "Email", "website": "Website", "address": "Address",
    "linkedin": "LinkedIn", "facebook": "Facebook", "instagram": "Instagram", "other_social": "Other social",
    "rating": "Rating", "phone_source": "Phone source", "website_source": "Website source",
    "decision_maker_source": "Decision-maker source", "source_url": "Listing", "map_url": "Map",
}


def to_export_frame(leads: list[dict] | pd.DataFrame) -> pd.DataFrame:
    df = pd.DataFrame(leads).copy()
    if df.empty:
        return pd.DataFrame(columns=["Readiness", *EXPORT_COLUMNS.values()])
    df["readiness"] = [readiness_score(r) for r in df.to_dict("records")]
    cols = ["readiness"] + [c for c in EXPORT_COLUMNS if c in df.columns]
    extra = [c for c in ("status", "notes") if c in df.columns]
    out = df[cols + extra].fillna("").rename(columns={**EXPORT_COLUMNS, "readiness": "Readiness", "status": "Status", "notes": "Notes"})
    return out


def to_csv_bytes(leads) -> bytes:
    return to_export_frame(leads).to_csv(index=False).encode("utf-8-sig")


def to_excel_bytes(leads) -> bytes:
    df = to_export_frame(leads)
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Leads")
        sheet = writer.sheets["Leads"]
        header_fill = PatternFill("solid", fgColor="0D2B2E")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for idx, column in enumerate(df.columns, 1):
            longest = max([len(str(column))] + [len(str(v)) for v in df[column].head(200)])
            sheet.column_dimensions[get_column_letter(idx)].width = min(max(11, longest + 2), 48)
        for idx, column in enumerate(df.columns, 1):
            if column in {"Website", "LinkedIn", "Facebook", "Instagram", "Listing", "Map"}:
                for row in range(2, sheet.max_row + 1):
                    cell = sheet.cell(row=row, column=idx)
                    if isinstance(cell.value, str) and cell.value.startswith("http"):
                        cell.hyperlink = cell.value.split(",")[0]
                        cell.font = Font(color="0E7C7B", underline="single")
    return buffer.getvalue()
